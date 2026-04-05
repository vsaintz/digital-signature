from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any


@dataclass(frozen=True)
class NormalizedDocument:
    columns: list[str]
    rows: list[list[Any]]
    version: int = 1

    def to_dict(self) -> dict:
        return {
            "version": self.version,
            "columns": self.columns,
            "rows": self.rows,
        }

    def to_json(self) -> str:
        return json.dumps(
            self.to_dict(),
            ensure_ascii=True,
            allow_nan=False,
            separators=(",", ":"),
        )


def _normalize_headers(raw_headers: list[str]) -> list[str]:
    normalized = [re.sub(r"\s+", " ", h.strip().lower()) for h in raw_headers]
    seen: dict[str, int] = {}
    result: list[str] = []
    for h in normalized:
        if h not in seen:
            seen[h] = 1
            result.append(h)
        else:
            seen[h] += 1
            result.append(f"{h}_{seen[h]}")
    return result


_INT_RE = re.compile(r"^-?\d+$")
_FLOAT_RE = re.compile(r"^-?\d+\.\d+$")
_BOOL_MAP = {"true": True, "false": False, "yes": True, "no": False}

_DATE_FORMATS = [
    ("%Y-%m-%d", False),
    ("%d/%m/%Y", False),
    ("%m/%d/%Y", False),
    ("%d-%m-%Y", False),
    ("%d %b %Y", False),
    ("%d %B %Y", False),
    ("%Y-%m-%dT%H:%M:%S", True),
    ("%Y-%m-%d %H:%M:%S", True),
    ("%d/%m/%Y %H:%M:%S", True),
]


def _normalize_value(raw: Any) -> Any:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return raw
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(raw, date):
        return raw.strftime("%Y-%m-%d")

    text = str(raw).strip()
    if not text:
        return None

    text = re.sub(r"\s+", " ", text)
    lower = text.lower()

    if lower in _BOOL_MAP:
        return _BOOL_MAP[lower]
    if _INT_RE.match(text):
        return int(text)
    if _FLOAT_RE.match(text):
        return float(text)

    for fmt, has_time in _DATE_FORMATS:
        try:
            dt = datetime.strptime(text, fmt)
            return (
                dt.strftime("%Y-%m-%dT%H:%M:%S")
                if has_time
                else dt.strftime("%Y-%m-%d")
            )
        except ValueError:
            continue

    return text


def _read_csv(file) -> tuple[list[str], list[list[Any]]]:
    file.seek(0)
    text = file.read().decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    file.seek(0)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _read_xlsx(file) -> tuple[list[str], list[list[Any]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for XLSX files. Run: pip install openpyxl"
        ) from exc

    file.seek(0)
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)

    ws = wb.active
    if ws is None:
        wb.close()
        file.seek(0)
        return [], []

    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    file.seek(0)

    if not rows:
        return [], []

    raw_headers = [str(c) if c is not None else "" for c in rows[0]]
    return raw_headers, [list(row) for row in rows[1:]]


_READERS = {
    "csv": _read_csv,
    "xlsx": _read_xlsx,
}


def normalize(file, file_type: str) -> NormalizedDocument:
    reader = _READERS.get(file_type)
    if reader is None:
        raise ValueError(f"No reader for file type '{file_type}'.")

    raw_headers, raw_rows = reader(file)
    columns = _normalize_headers(raw_headers)
    n_cols = len(columns)

    rows = []
    for raw_row in raw_rows:
        padded = list(raw_row) + [None] * max(0, n_cols - len(raw_row))
        normalized_row = [_normalize_value(cell) for cell in padded[:n_cols]]
        if any(v is not None for v in normalized_row):
            rows.append(normalized_row)

    return NormalizedDocument(columns=columns, rows=rows)
